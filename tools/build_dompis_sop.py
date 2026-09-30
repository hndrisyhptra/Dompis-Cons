from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Claude outputs" / "SOP_BUKU_SAKU_FLOW_DOMPIS_CONS_PT3_PT2_2026-09-28.md"
OUTPUT = ROOT / "Claude outputs" / "SOP_BUKU_SAKU_FLOW_DOMPIS_CONS_PT3_PT2_2026-09-28.docx"

NAVY = "163A5F"
BLUE = "1565D8"
LIGHT_BLUE = "EAF2FB"
PALE_BLUE = "F5F8FC"
PALE_GRAY = "F6F7F9"
MID_GRAY = "6B7280"
BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)
    keep_next.set(qn("w:val"), "1" if value else "0")


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Aptos", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text: str, *, size=None, color=BLACK) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Aptos Mono", size=(size or 10.5) - 0.5, color="334155")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_run_font(run, size=8.5, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc: Document) -> None:
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.30)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 17, 14, 7),
        ("Heading 2", 13.5, 11, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.3)
        style.paragraph_format.left_indent = Inches(0.26)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(3.5)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("DOMPIS CONS   |   SOP OPERASIONAL PT3 DAN PT2")
    set_run_font(hr, size=8.2, color=MID_GRAY, bold=True)

    footer = section.footer
    ft = footer.add_table(rows=1, cols=2, width=Inches(7.0))
    ft.alignment = WD_TABLE_ALIGNMENT.CENTER
    ft.columns[0].width = Inches(4.8)
    ft.columns[1].width = Inches(2.2)
    left = ft.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lr = left.add_run("Versi 1.0   |   Draf untuk validasi bisnis   |   28 September 2026")
    set_run_font(lr, size=8.0, color=MID_GRAY)
    add_page_number(ft.cell(0, 1).paragraphs[0])


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(60)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("DOMPIS CONS")
    set_run_font(run, name="Aptos Display", size=13, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(10)
    title.add_run("SOP dan Buku Saku Alur Operasional")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(32)
    sr = subtitle.add_run("Project PT3 Reguler dan PT2")
    set_run_font(sr, name="Aptos Display", size=18, color=NAVY, bold=True)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(28)
    add_inline(
        intro,
        "Panduan kerja end to end sejak data PID dan BOQ disiapkan sampai setiap LOP diverifikasi dan berstatus Golive.",
        size=12,
        color="334155",
    )

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.65)
    table.columns[1].width = Inches(4.55)
    meta = [
        ("Versi", "1.0"),
        ("Tanggal", "28 September 2026"),
        ("Status", "Draf untuk validasi bisnis"),
        ("Lingkup", "Operasional PT3 Reguler dan PT2"),
    ]
    for idx, (label, value) in enumerate(meta):
        for cell in table.rows[idx].cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=120, bottom=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(table.cell(idx, 0), LIGHT_BLUE)
        p1 = table.cell(idx, 0).paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(label)
        set_run_font(r1, size=9.5, color=NAVY, bold=True)
        p2 = table.cell(idx, 1).paragraphs[0]
        r2 = p2.add_run(value)
        set_run_font(r2, size=9.8, color=BLACK)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(46)
    nr = note.add_run("Dokumen kerja internal")
    set_run_font(nr, size=9, color=MID_GRAY, italic=True)
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    h = doc.add_paragraph("Daftar isi ringkas", style="Heading 1")
    h.paragraph_format.space_before = Pt(0)
    intro = doc.add_paragraph(
        "Gunakan daftar ini untuk menemukan flow, gate approval, dan kontrol harian yang dibutuhkan."
    )
    intro.paragraph_format.space_after = Pt(14)
    sections = [
        ("1", "Tujuan dan cara menggunakan buku saku"),
        ("2", "Prinsip operasional utama"),
        ("3", "Peran dan tanggung jawab"),
        ("4", "Flow ringkas end to end"),
        ("5", "SOP end to end PT3 Reguler"),
        ("6", "SOP end to end PT2"),
        ("7", "Aturan approval dan koreksi"),
        ("8", "Alur pengecualian"),
        ("9", "Monitoring dan pelaporan"),
        ("10", "Checklist harian per peran"),
        ("11", "Referensi status cepat"),
        ("12", "Matriks gate akhir"),
        ("13", "Eskalasi dan kontrol perubahan"),
    ]
    table = doc.add_table(rows=len(sections), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(0.55)
    table.columns[1].width = Inches(5.9)
    for idx, (number, label) in enumerate(sections):
        row = table.rows[idx]
        set_cant_split(row)
        for cell in row.cells:
            set_cell_border(cell, color=WHITE, size=0)
            set_cell_margins(cell, top=60, bottom=60)
            if idx % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
        p1 = row.cells[0].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(number)
        set_run_font(r1, size=10, color=BLUE, bold=True)
        p2 = row.cells[1].paragraphs[0]
        r2 = p2.add_run(label)
        set_run_font(r2, size=10, color=BLACK, bold=(number in {"5", "6"}))
    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        set_cant_split(row)
        for c_idx, value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            if c_idx == 0 and value.isdigit():
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(
                paragraph,
                value,
                size=8.7 if col_count >= 4 else 9.1,
                color=WHITE if r_idx == 0 else BLACK,
            )
            if r_idx == 0:
                for run in paragraph.runs:
                    run.bold = True

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def is_table_separator(line: str) -> bool:
    stripped = line.strip().strip("|")
    parts = [part.strip() for part in stripped.split("|")]
    return bool(parts) and all(re.fullmatch(r":?-{3,}:?", part) for part in parts)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        line = lines[idx]
        if not is_table_separator(line):
            rows.append([part.strip() for part in line.strip().strip("|").split("|")])
        idx += 1
    return rows, idx


def add_flow_paragraph(doc: Document, text: str) -> None:
    steps = [part.strip() for part in text.split("→")]
    table = doc.add_table(rows=len(steps), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(0.55)
    table.columns[1].width = Inches(5.95)
    for idx, step in enumerate(steps):
        row = table.rows[idx]
        set_cant_split(row)
        for cell in row.cells:
            set_cell_border(cell, color=WHITE, size=0)
            set_cell_margins(cell, top=50, bottom=50)
            if idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
        p_num = row.cells[0].paragraphs[0]
        p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rn = p_num.add_run(str(idx + 1))
        set_run_font(rn, size=9, color=BLUE, bold=True)
        p_step = row.cells[1].paragraphs[0]
        add_inline(p_step, step.replace("**", ""), size=9.7, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_body_from_markdown(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1 "))
    lines = lines[start:]
    page_break_before = {
        "4.2 PT2",
        "5 SOP end to end PT3 Reguler",
        "10 Checklist harian per peran",
        "11 Referensi status cepat",
    }

    idx = 0
    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()
        if not line:
            idx += 1
            continue

        if line.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading_match:
            hashes, heading = heading_match.groups()
            if heading in page_break_before:
                doc.add_page_break()
            level = len(hashes) - 1
            paragraph = doc.add_paragraph(heading, style=f"Heading {level}")
            if level == 1 and heading.startswith(("1 ", "5 ", "6 ", "9 ", "10 ", "11 ")):
                paragraph.paragraph_format.space_before = Pt(0)
            idx += 1
            continue

        if "→" in line and line.startswith("**"):
            add_flow_paragraph(doc, line)
            idx += 1
            continue

        number_match = re.match(r"^(\d+)\.\s+(.+)$", line)
        if number_match:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(3.5)
            paragraph.paragraph_format.line_spacing = 1.08
            number_run = paragraph.add_run(number_match.group(1) + ".  ")
            set_run_font(number_run, size=10.3, color=BLUE, bold=True)
            add_inline(paragraph, number_match.group(2), size=10.3)
            idx += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", line)
        if bullet_match:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.20)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(3.5)
            paragraph.paragraph_format.line_spacing = 1.08
            bullet_run = paragraph.add_run("•  ")
            set_run_font(bullet_run, size=10.3, color=BLUE, bold=True)
            add_inline(paragraph, bullet_match.group(1), size=10.3)
            idx += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(1)
        add_inline(paragraph, line, size=10.5)
        if line.startswith(("**Pelaksana:**", "**Reviewer:**", "**Langkah kerja:**", "**Gate dan hasil:**", "**Hasil:**", "**Kontrol:**", "**Acuan item:**", "**Jalur dokumen pendukung:**")):
            paragraph.paragraph_format.space_after = Pt(5)
        idx += 1


def add_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "SOP dan Buku Saku Alur Operasional DOMPIS CONS"
    props.subject = "Flow end to end Project PT3 Reguler dan PT2"
    props.author = "DOMPIS CONS"
    props.keywords = "DOMPIS CONS, SOP, PT3, PT2, LOP, approval, Golive"
    props.comments = "Versi 1.0 tanggal 28 September 2026"


def build() -> None:
    doc = Document()
    configure_document(doc)
    add_document_properties(doc)
    add_cover(doc)
    add_contents(doc)
    add_body_from_markdown(doc)

    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            set_keep_with_next(paragraph)
        paragraph.paragraph_format.widow_control = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
